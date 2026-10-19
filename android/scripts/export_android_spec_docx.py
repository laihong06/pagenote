#!/usr/bin/env python3
"""Export the Android requirements, design, and task ledger to one DOCX."""

from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ANDROID_ROOT = Path(__file__).resolve().parent.parent
SPEC_ROOT = ANDROID_ROOT / "docs" / "spec"
OUTPUT = ANDROID_ROOT / "docs" / "Android-SPEC.docx"
FONT = "Hiragino Sans GB"
MONO_FONT = "Menlo"
TICK = chr(96)


def set_run_font(run, name=FONT, size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run_properties = run._element.get_or_add_rPr()
    fonts = run_properties.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        run_properties.insert(0, fonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_inline(paragraph, text):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    inline_code = re.escape(TICK) + r"[^" + re.escape(TICK) + r"]+" + re.escape(TICK)
    token_pattern = re.compile(
        r"(\*\*.+?\*\*|"
        + inline_code
        + r"|\[[^\]]+\]\([^)]+\)|\*[^*]+\*|(?<!\w)_(?!_)[^_]+(?<!_)_(?!\w))"
    )
    cursor = 0
    for match in token_pattern.finditer(text):
        if match.start() > cursor:
            set_run_font(paragraph.add_run(text[cursor : match.start()]))
        token = match.group(0)
        if token.startswith("**"):
            set_run_font(paragraph.add_run(token[2:-2]), bold=True)
        elif token.startswith(TICK):
            set_run_font(paragraph.add_run(token[1:-1]), name=MONO_FONT, size=9)
        elif token.startswith("["):
            link = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", token)
            if link:
                set_run_font(paragraph.add_run(link.group(1)), color="185ABD")
                set_run_font(paragraph.add_run(f" ({link.group(2)})"), size=9)
            else:
                set_run_font(paragraph.add_run(token))
        elif token.startswith("_"):
            set_run_font(paragraph.add_run(token[1:-1]), italic=True)
        else:
            set_run_font(paragraph.add_run(token[1:-1]), italic=True)
        cursor = match.end()
    if cursor < len(text):
        set_run_font(paragraph.add_run(text[cursor:]))


def add_table(document, rows):
    cells = [[cell.strip() for cell in row.strip().strip("|").split("|")] for row in rows]
    cells = [row for row in cells if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in row)]
    if not cells:
        return
    column_count = max(len(row) for row in cells)
    table = document.add_table(rows=1, cols=column_count)
    table.style = "Table Grid"
    for row_index, values in enumerate(cells):
        target = table.rows[0].cells if row_index == 0 else table.add_row().cells
        for column_index in range(column_count):
            paragraph = target[column_index].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            add_inline(paragraph, values[column_index] if column_index < len(values) else "")
            if row_index == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "185ABD")
                target[column_index]._tc.get_or_add_tcPr().append(shading)
                for run in paragraph.runs:
                    set_run_font(run, color="FFFFFF", bold=True, size=9)
            else:
                for run in paragraph.runs:
                    if run.font.size is None:
                        set_run_font(run, size=9)


def add_markdown_source(document, source_path, section_title):
    document.add_page_break()
    document.add_heading(section_title, level=1)
    lines = source_path.read_text(encoding="utf-8").splitlines()
    in_code = False
    code_lines = []
    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if stripped.startswith(TICK * 3):
            if in_code:
                paragraph = document.add_paragraph(style="Code Block")
                set_run_font(paragraph.add_run("\n".join(code_lines)), name=MONO_FONT, size=8.5)
                in_code = False
                code_lines = []
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if not stripped or stripped == "---":
            index += 1
            continue
        if stripped.startswith("# "):
            index += 1
            continue
        heading = re.match(r"^(#{2,4})\s+(.+)$", stripped)
        if heading:
            level = min(len(heading.group(1)), 3)
            paragraph = document.add_paragraph(style=f"Heading {level}")
            add_inline(paragraph, heading.group(2))
            index += 1
            continue
        if stripped.startswith("|"):
            table_rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_rows.append(lines[index].strip())
                index += 1
            add_table(document, table_rows)
            continue
        if stripped.startswith(">"):
            paragraph = document.add_paragraph(style="Quote")
            add_inline(paragraph, stripped.lstrip("> "))
            index += 1
            continue
        listing = re.match(r"^(\s*)(?:[-+*]|\d+\.)\s+(.+)$", line)
        if listing:
            paragraph = document.add_paragraph(
                style="List Number" if re.match(r"^\s*\d+\.", line) else "List Bullet"
            )
            paragraph.paragraph_format.left_indent = Inches(min(len(listing.group(1)) // 2, 4) * 0.2)
            add_inline(paragraph, listing.group(2))
            index += 1
            continue
        paragraph = document.add_paragraph()
        add_inline(paragraph, stripped)
        index += 1
    if in_code and code_lines:
        paragraph = document.add_paragraph(style="Code Block")
        set_run_font(paragraph.add_run("\n".join(code_lines)), name=MONO_FONT, size=8.5)


def configure_document(document):
    section = document.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(9)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.0

    title = document.styles["Title"]
    title.font.name = FONT
    title.font.size = Pt(27)
    title.font.bold = True
    title.font.underline = False
    title.font.color.rgb = RGBColor.from_string("253247")
    title.paragraph_format.space_after = Pt(8)
    title_properties = title._element.pPr
    if title_properties is not None:
        for border in title_properties.findall(qn("w:pBdr")):
            title_properties.remove(border)

    for style_name, size, color in (
        ("Heading 1", 17, "185ABD"),
        ("Heading 2", 13, "253247"),
        ("Heading 3", 11, "536276"),
    ):
        style = document.styles[style_name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    code = document.styles.add_style("Code Block", 1)
    code.font.name = MONO_FONT
    code.font.size = Pt(8.5)
    code.paragraph_format.left_indent = Inches(0.14)
    code.paragraph_format.right_indent = Inches(0.08)
    code.paragraph_format.space_before = Pt(3)
    code.paragraph_format.space_after = Pt(6)
    code.paragraph_format.keep_together = True
    quote = document.styles["Quote"]
    quote.font.name = FONT
    quote.font.size = Pt(9)
    quote.font.color.rgb = RGBColor.from_string("536276")
    quote.paragraph_format.left_indent = Inches(0.22)


def main():
    document = Document()
    configure_document(document)
    document.core_properties.title = "Pagenote Android SPEC"
    document.core_properties.subject = "Android product requirements, design, and task status"
    document.core_properties.keywords = "pagenote, Android, SPEC"

    title = document.add_paragraph(style="Title")
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_properties = title._p.get_or_add_pPr()
    for border in title_properties.findall(qn("w:pBdr")):
        title_properties.remove(border)
    add_inline(title, "Pagenote Android SPEC")
    subtitle = document.add_paragraph()
    set_run_font(
        subtitle.add_run("原生 Android 需求、技术设计与交付状态"),
        size=12,
        color="536276",
    )
    snapshot = document.add_paragraph()
    set_run_font(snapshot.add_run("状态快照：2026-10-18"), size=9, color="536276")
    intro = document.add_paragraph()
    add_inline(
        intro,
        "本文汇总 Android 端需求、技术设计和任务状态。具体功能契约仍以 android/docs/spec 中的 Markdown 文档为实施依据。"
    )
    document.add_heading("当前验证结果", level=1)
    for item in (
        "Android 14 4 KB、Android 15 4 KB 与 Android 15 16 KB 三种 AVD 各通过 107 项 app 仪器测试。",
        "app 仪器测试 107/107、原生手写笔 View 仪器测试 3/3 均在上述三种 AVD 各通过；实体笔硬件验收仍待完成。",
        "Android JVM 测试合计通过 370 项；assembleDebug 构建成功。",
        "四 ABI debug APK 均通过 16 KB ELF LOAD 与 ZIP 对齐检查。",
        "task 12 仍未关闭：完整手动工作流、实体手写笔、自然低内存进程恢复和真实 OpenAI 兼容服务仍待验证。"
    ):
        paragraph = document.add_paragraph(style="List Bullet")
        add_inline(paragraph, item)

    add_markdown_source(document, SPEC_ROOT / "requirements.md", "Android 端需求说明")
    add_markdown_source(document, SPEC_ROOT / "design.md", "Android 端设计说明")
    add_markdown_source(document, SPEC_ROOT / "tasks.md", "Android 任务状态")
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
